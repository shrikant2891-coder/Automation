#!/usr/bin/env python3
"""Structural checks for upgraded payroll workbook (no Excel runtime required)."""

from pathlib import Path

from openpyxl import load_workbook

PATH = Path("/workspace/Automatic_Salary_Slip_PF_IT_Fixed.xlsm")
REQUIRED_SHEETS = [
    "Generator",
    "Payroll Settings",
    "Salary sheet",
    "Income Tax",
    "Incremental sheet",
    "Salary Slip",
    "Master Data",
]


def main() -> None:
    assert PATH.is_file(), f"Missing {PATH}"
    wb = load_workbook(PATH, keep_vba=True)
    for name in REQUIRED_SHEETS:
        assert name in wb.sheetnames, f"Missing sheet: {name}"

    settings = wb["Payroll Settings"]
    assert settings["B2"].value == 1800
    assert settings["B3"].value == 3000
    assert settings["B4"].value == 6

    sal = wb["Salary sheet"]
    assert "Sep+" in str(sal["BT2"].value)
    assert sal["H3"].value == '=IF(D3="Yes",\'Payroll Settings\'!$B$2,0)'
    assert sal["I3"].value == '=IF(D3="Yes",\'Payroll Settings\'!$B$3,0)'
    assert "BU3" in sal["AR3"].value or "BU3" in str(sal["AR3"].value)

    slip = wb["Salary Slip"]
    assert "BT$3" in slip["M14"].value
    assert "Payroll Settings" in slip["I21"].value

    import zipfile

    with zipfile.ZipFile(PATH) as z:
        assert "xl/vbaProject.bin" in z.namelist()

    print("OK: payroll workbook structure validated")


if __name__ == "__main__":
    main()
