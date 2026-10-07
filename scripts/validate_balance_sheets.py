#!/usr/bin/env python3
"""Validate that balance sheet sources equal application of funds in FS workbooks."""

from __future__ import annotations

import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("Install openpyxl: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]

WORKBOOKS = [
    {
        "path": "WEBKEY_SOLUTIONS_FS_FY2025-26_Comparable.xlsx",
        "sheet": "Balance Sheet",
        "liab_row": 24,
        "asset_row": 34,
        "col": 4,
    },
    {
        "path": "Gupta_Agency_Financial_Statements_FY2025-26.xlsx",
        "sheet": "Balance Sheet",
        "liab_row": 18,
        "asset_row": 31,
        "col": 3,
    },
    {
        "path": "Shree_Ganeshay_Traders_Financial_Statements_FY2025-26.xlsx",
        "sheet": "Balance Sheet",
        "liab_row": 17,
        "asset_row": 28,
        "col": 3,
    },
    {
        "path": "Shree_Balajee_Vertical_FS_25-26.xlsx",
        "sheet": "Balance Sheet",
        "liab_row": 24,
        "asset_row": 37,
        "col": 3,
    },
    {
        "path": "HNSPVTL_FS_FY2025-26_Final.xlsx",
        "sheet": "Balance Sheet",
        "liab_row": 25,
        "asset_row": 34,
        "col": 4,
    },
]

TOLERANCE = 0.02  # rupees / hundreds


def main() -> int:
    failed = 0
    for spec in WORKBOOKS:
        path = ROOT / spec["path"]
        if not path.exists():
            print(f"SKIP  {spec['path']} (file not found)")
            continue
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb[spec["sheet"]]
        c = spec["col"]
        liab = ws.cell(spec["liab_row"], c).value
        assets = ws.cell(spec["asset_row"], c).value
        wb.close()
        if liab is None or assets is None:
            print(f"FAIL  {spec['path']}: missing total cells")
            failed += 1
            continue
        diff = float(liab) - float(assets)
        status = "OK" if abs(diff) <= TOLERANCE else "FAIL"
        print(f"{status}  {spec['path']}: liabilities={liab}, assets={assets}, diff={diff}")
        if status == "FAIL":
            failed += 1
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
