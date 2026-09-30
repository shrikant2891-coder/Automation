#!/usr/bin/env python3
"""Export Automatic_Salary_Slip_PF_IT_Fixed.xlsm into payroll-dashboard/config/payroll.xml."""

from __future__ import annotations

import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

XLSM = Path("/workspace/Automatic_Salary_Slip_PF_IT_Fixed.xlsm")
OUT = Path("/workspace/payroll-dashboard/config/payroll.xml")

MONTHS = [
    ("1", "Apr'26", 30),
    ("2", "May'26", 31),
    ("3", "Jun'26", 30),
    ("4", "Jul'26", 31),
    ("5", "Aug'26", 31),
    ("6", "Sep'26", 30),
    ("7", "Oct'26", 31),
    ("8", "Nov'26", 30),
    ("9", "Dec'26", 31),
    ("10", "Jan'27", 31),
    ("11", "Feb'27", 28),
    ("12", "Mar'27", 31),
]
ATT_COLS = list("LMNOPQRSTUVW")


def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip()


def main() -> None:
    wb = load_workbook(XLSM, data_only=True)
    wb_f = load_workbook(XLSM, data_only=False)
    root = ET.Element(
        "payroll",
        {
            "fy": "2026-27",
            "version": "1",
            "updated": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        },
    )

    settings = ET.SubElement(root, "settings")
    ps = wb["Payroll Settings"] if "Payroll Settings" in wb.sheetnames else None
    ET.SubElement(
        settings,
        "setting",
        name="pf_cap_legacy",
        value=str(int(ps["B2"].value or 1800) if ps else 1800),
    )
    ET.SubElement(
        settings,
        "setting",
        name="pf_cap_new",
        value=str(int(ps["B3"].value or 3000) if ps else 3000),
    )
    ET.SubElement(
        settings,
        "setting",
        name="pf_effective_month_index",
        value=str(int(ps["B4"].value or 6) if ps else 6),
    )
    ET.SubElement(
        settings,
        "setting",
        name="pf_effective_month_label",
        value=text(ps["B5"].value if ps else "Sep'26") or "Sep'26",
    )
    ET.SubElement(settings, "setting", name="output_folder", value="")

    months_el = ET.SubElement(root, "months")
    for idx, label, days in MONTHS:
        ET.SubElement(months_el, "month", index=idx, label=label, days=str(days))

    master_ws = wb["Master Data"]
    master_f = wb_f["Master Data"]
    salary_ws = wb["Salary sheet"]
    salary_f = wb_f["Salary sheet"]
    inc_ws = wb["Incremental sheet"]

    employees_el = ET.SubElement(root, "employees")
    for r in range(3, 500):
        code = salary_ws[f"A{r}"].value
        if not code:
            break
        emp = ET.SubElement(employees_el, "employee", code=text(code))
        ET.SubElement(emp, "salary", basic=text(salary_ws[f"E{r}"].value))
        ET.SubElement(emp, "salary", pf_status=text(salary_ws[f"D{r}"].value) or "No")
        ET.SubElement(emp, "salary", status=text(salary_ws[f"C{r}"].value) or "Active")

        att_el = ET.SubElement(emp, "attendance")
        for i, col in enumerate(ATT_COLS):
            days = salary_ws[f"{col}{r}"].value
            if days is None or days == "":
                continue
            ET.SubElement(
                att_el,
                "month",
                index=str(i + 1),
                label=MONTHS[i][1],
                days=text(days),
            )

        mrow = None
        for mr in range(2, 500):
            if text(master_ws[f"A{mr}"].value) == text(code):
                mrow = mr
                break
        master_el = ET.SubElement(emp, "master")
        if mrow:
            fields = [
                ("name", "B"),
                ("father_name", "C"),
                ("email", "D"),
                ("dob", "E"),
                ("gender", "F"),
                ("doj", "G"),
                ("uan", "H"),
                ("pf_account", "I"),
                ("pan", "J"),
                ("bank_code", "K"),
                ("bank_name", "L"),
                ("bank_account", "M"),
                ("ifsc", "N"),
                ("department", "O"),
                ("designation", "P"),
                ("group", "Q"),
            ]
            for tag, col in fields:
                val = master_ws[f"{col}{mrow}"].value
                if val not in (None, ""):
                    ET.SubElement(master_el, tag).text = text(val)

    increments_el = ET.SubElement(root, "increments")
    for r in range(9, 200):
        code = inc_ws[f"A{r}"].value
        if not code or not str(code).startswith("INF"):
            continue
        ET.SubElement(
            increments_el,
            "increment",
            emp_code=text(code),
            effective_month=text(inc_ws[f"B{r}"].value),
            new_basic=text(inc_ws[f"C{r}"].value),
            old_basic=text(inc_ws[f"D{r}"].value),
            arrear_months=text(inc_ws[f"E{r}"].value or "3"),
            notes=text(inc_ws[f"F{r}"].value),
        )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    tree = ET.ElementTree(root)
    ET.indent(tree, space="  ")
    tree.write(OUT, encoding="unicode", xml_declaration=True)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
